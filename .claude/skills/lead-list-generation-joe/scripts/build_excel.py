#!/usr/bin/env python3
"""Schritt 6: alles zu einer Excel-Datei zusammenfuehren.

Nimmt die gefilterten Personen, die BetterContact-Ergebnisse und die Lusha-Ergebnisse
und baut daraus output/06_kaeuferliste.xlsx mit drei Blaettern:

  Anrufliste    eine Zeile je Person, Anrufbare zuerst
  Firmen        die Firmen aus Schritt 1
  Zusammenfassung  was reinging, was rauskam, wo es verloren ging

Die beiden Anreicherungs-Quellen stehen in GETRENNTEN Spalten. Zusaetzlich gibt es
je eine Spalte "beste E-Mail" und "beste Nummer" zum direkten Arbeiten. Getrennt
bleiben sie, damit man bei einer falschen Nummer weiss, wer sie geliefert hat.
"""
import argparse
import csv
import sys

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from _common import OUT, die, read_json  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    die("Das Paket openpyxl fehlt. Einmalig installieren:\n\n    pip3 install openpyxl\n")

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from lusha_enrich import treffer_aus  # noqa: E402

SPALTEN = [
    ("person_id", 10), ("vorname", 14), ("nachname", 16), ("titel", 30),
    ("firma", 26), ("firma_domain", 22), ("ort", 18),
    ("passung_score", 8), ("passung_begruendung", 46),
    ("beste_email", 30), ("beste_nummer", 20), ("anrufbar", 10),
    ("email_bettercontact", 30), ("nummer_bettercontact", 20),
    ("email_lusha", 30), ("nummer_lusha", 20),
    ("linkedin_url", 40), ("firma_linkedin", 40),
]


def csv_lesen(pfad, pflicht=True):
    try:
        with open(pfad, newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    except FileNotFoundError:
        if pflicht:
            die(f"{pfad} fehlt.")
        return []


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--personen", default=str(OUT / "03_personen_gefiltert.csv"))
    ap.add_argument("--firmen", default=str(OUT / "01_firmen.csv"))
    ap.add_argument("--ziel", default=str(OUT / "06_kaeuferliste.xlsx"))
    a = ap.parse_args()

    personen = csv_lesen(a.personen)
    firmen = csv_lesen(a.firmen, pflicht=False)
    if not personen:
        die("Keine Personen. Erst Schritt 3 abschliessen.")

    bc = {}
    for zeile in read_json(OUT / "04_bettercontact_raw.json", []):
        pid = (zeile.get("custom_fields") or {}).get("person_id")
        if pid:
            bc[pid] = zeile
    lusha = read_json(OUT / "05_lusha_raw.json", {})

    zeilen = []
    for p in personen:
        pid = p.get("person_id", "")
        b = bc.get(pid, {})
        bc_mail = (b.get("contact_email_address") or "").strip()
        bc_tel = (b.get("contact_phone_number") or "").strip()
        l_mails, l_tels = treffer_aus(lusha.get(pid, {}).get("antwort", {}))
        l_mail = l_mails[0] if l_mails else ""
        l_tel = l_tels[0] if l_tels else ""

        beste_mail = bc_mail or l_mail
        beste_tel = bc_tel or l_tel
        zeilen.append({
            "person_id": pid,
            "vorname": p.get("vorname", ""),
            "nachname": p.get("nachname", ""),
            "titel": p.get("titel", ""),
            "firma": p.get("firma", ""),
            "firma_domain": p.get("firma_domain", ""),
            "ort": p.get("ort", ""),
            "passung_score": p.get("passung_score", ""),
            "passung_begruendung": p.get("passung_begruendung", ""),
            "beste_email": beste_mail,
            "beste_nummer": beste_tel,
            "anrufbar": "ja" if beste_tel else "nein",
            "email_bettercontact": bc_mail,
            "nummer_bettercontact": bc_tel,
            "email_lusha": l_mail,
            "nummer_lusha": l_tel,
            "linkedin_url": p.get("linkedin_url", ""),
            "firma_linkedin": p.get("firma_linkedin", ""),
        })

    def sortierschluessel(z):
        try:
            score = -float(z.get("passung_score") or 0)
        except ValueError:
            score = 0
        return (0 if z["beste_nummer"] else 1, 0 if z["beste_email"] else 1, score)

    zeilen.sort(key=sortierschluessel)

    wb = Workbook()
    kopf_fuellung = PatternFill("solid", fgColor="1F2937")
    kopf_schrift = Font(color="FFFFFF", bold=True)

    ws = wb.active
    ws.title = "Anrufliste"
    ws.append([s for s, _ in SPALTEN])
    for i, (name, breite) in enumerate(SPALTEN, 1):
        z = ws.cell(row=1, column=i)
        z.fill = kopf_fuellung
        z.font = kopf_schrift
        z.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = breite
    for z in zeilen:
        ws.append([z.get(s, "") for s, _ in SPALTEN])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{len(zeilen) + 1}"

    wsf = wb.create_sheet("Firmen")
    if firmen:
        kopf = list(firmen[0].keys())
        wsf.append(kopf)
        for i in range(1, len(kopf) + 1):
            zelle = wsf.cell(row=1, column=i)
            zelle.fill = kopf_fuellung
            zelle.font = kopf_schrift
            wsf.column_dimensions[get_column_letter(i)].width = 30
        for f in firmen:
            wsf.append([f.get(k, "") for k in kopf])
        wsf.freeze_panes = "A2"

    mit_tel = sum(1 for z in zeilen if z["beste_nummer"])
    mit_mail = sum(1 for z in zeilen if z["beste_email"])
    beides = sum(1 for z in zeilen if z["beste_nummer"] and z["beste_email"])
    nur_bc_tel = sum(1 for z in zeilen if z["nummer_bettercontact"] and not z["nummer_lusha"])
    nur_l_tel = sum(1 for z in zeilen if z["nummer_lusha"] and not z["nummer_bettercontact"])

    wsz = wb.create_sheet("Zusammenfassung")
    for a_, b_ in [
        ("Was", "Anzahl"),
        ("Firmen in Schritt 1", len(firmen)),
        ("Personen nach der Passungs-Pruefung", len(zeilen)),
        ("davon mit Mobilnummer", mit_tel),
        ("davon mit E-Mail", mit_mail),
        ("davon mit beidem", beides),
        ("Nummer nur von BetterContact", nur_bc_tel),
        ("Nummer nur von Lusha", nur_l_tel),
        ("ohne jede Kontaktangabe", len(zeilen) - mit_tel - mit_mail + beides),
    ]:
        wsz.append([a_, b_])
    wsz.cell(row=1, column=1).font = kopf_schrift
    wsz.cell(row=1, column=1).fill = kopf_fuellung
    wsz.cell(row=1, column=2).font = kopf_schrift
    wsz.cell(row=1, column=2).fill = kopf_fuellung
    wsz.column_dimensions["A"].width = 40
    wsz.column_dimensions["B"].width = 12

    wb.save(a.ziel)
    print(f"\nFertig: {a.ziel}")
    print(f"  Personen:            {len(zeilen)}")
    print(f"  mit Mobilnummer:     {mit_tel}")
    print(f"  mit E-Mail:          {mit_mail}")
    print(f"  Nummer nur von Lusha:        {nur_l_tel}")
    print(f"  Nummer nur von BetterContact: {nur_bc_tel}")
    if mit_tel == 0:
        print("\n  ACHTUNG: keine einzige Nummer. Das ist fast nie richtig - pruefe, ob die "
              "Anreicherungs-Schritte wirklich durchgelaufen sind.")


if __name__ == "__main__":
    main()
